import streamlit as st
import firebase_admin
from firebase_admin import credentials, firestore, auth, storage
import PyPDF2
import docx
import io
from io import BytesIO
from PIL import Image, ImageDraw
import pytesseract
from langchain.prompts import ChatPromptTemplate
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.schema.output_parser import StrOutputParser
from langchain.schema.runnable import RunnablePassthrough
from langchain.memory import ConversationBufferMemory
import base64
from langchain.text_splitter import RecursiveCharacterTextSplitter
import time
import random
from transformers import pipeline
import json

# Initialize Firebase Admin
# Initialize Firebase Admin
@st.cache_resource
def init_firebase_admin():
    if not firebase_admin._apps:
        try:
            # Create the credentials dictionary from the TOML format
            firebase_creds = {
                "type": st.secrets["firebase_credentials"]["type"],
                "project_id": st.secrets["firebase_credentials"]["project_id"],
                "private_key_id": st.secrets["firebase_credentials"]["private_key_id"],
                "private_key": st.secrets["firebase_credentials"]["private_key"],
                "client_email": st.secrets["firebase_credentials"]["client_email"],
                "client_id": st.secrets["firebase_credentials"]["client_id"],
                "auth_uri": st.secrets["firebase_credentials"]["auth_uri"],
                "token_uri": st.secrets["firebase_credentials"]["token_uri"],
                "auth_provider_x509_cert_url": st.secrets["firebase_credentials"]["auth_provider_x509_cert_url"],
                "client_x509_cert_url": st.secrets["firebase_credentials"]["client_x509_cert_url"],
                "universe_domain": st.secrets["firebase_credentials"]["universe_domain"]
            }
            
            cred = credentials.Certificate(firebase_creds)
            firebase_admin.initialize_app(cred, {
                'storageBucket': f"{st.secrets['firebase_config']['project_id']}.appspot.com"
            })
        except Exception as e:
            st.error(f"Error initializing Firebase: {str(e)}")
    return firestore.client()

# Initialize the LLM with secure API key
@st.cache_resource
def get_llm():
    return ChatGoogleGenerativeAI(
        model='gemini-pro',
        google_api_key=st.secrets["credentials"]["google_api_key"]
    )

def extract_text_from_file(file):
    """Extract text from various file types"""
    if file.type == "application/pdf":
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() or ""
    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(file)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    elif file.type.startswith('image'):
        image = Image.open(file)
        if image.mode != 'RGB':
            image = image.convert('RGB')
        text = pytesseract.image_to_string(image)
    else:
        text = file.getvalue().decode('utf-8', errors='ignore')
    return text

def generate_thumbnail(file):
    """Generate thumbnail for uploaded files"""
    if file.type.startswith('image'):
        image = Image.open(file)
        image.thumbnail((200, 200))
        buffered = BytesIO()
        image.save(buffered, format="PNG")
        return buffered.getvalue()
    else:
        img = Image.new('RGB', (200, 200), color='white')
        d = ImageDraw.Draw(img)
        d.text((10, 10), f"{file.name}\n\nDocument uploaded", fill=(0, 0, 0))
        buffered = BytesIO()
        img.save(buffered, format='PNG')
        return buffered.getvalue()

def chunk_text(text, chunk_size=1000, chunk_overlap=200):
    """Split text into manageable chunks"""
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len
    )
    return text_splitter.split_text(text)

def process_with_retry(func, *args, max_retries=3, **kwargs):
    """Process text with retry mechanism"""
    for attempt in range(max_retries):
        try:
            return func(*args, **kwargs)
        except Exception as e:
            if "429" in str(e) and attempt < max_retries - 1:
                wait_time = (2 ** attempt) + random.random()
                time.sleep(wait_time)
            elif attempt == max_retries - 1:
                st.warning("API limit reached. Falling back to local model.")
                fallback_model = get_fallback_model()
                return fallback_model(*args, max_length=100)[0]['generated_text']
            else:
                raise e

@st.cache_resource
def get_fallback_model():
    """Initialize fallback model"""
    return pipeline("text2text-generation", model="google/flan-t5-base", device="cpu")

def upload_file_to_firebase(user_id, file):
    """Upload file to Firebase Storage"""
    try:
        bucket = storage.bucket()
        blob_name = f"documents/{user_id}/{datetime.now().strftime('%Y%m%d_%H%M%S')}_{file.name}"
        blob = bucket.blob(blob_name)
        
        blob.upload_from_string(
            file.getvalue(),
            content_type=file.type
        )
        
        blob.make_public()
        return blob.public_url
    except Exception as e:
        st.error(f"Error uploading file: {str(e)}")
        return None

def generate_document_summary(document_chunks, llm):
    """Generate summary of document chunks"""
    summary_prompt = ChatPromptTemplate.from_template("""
    Please provide a concise summary of the following document chunk. 
    Focus on the main topics, key points, and important details:

    {doc_chunk}

    Summary:
    """)

    summary_chain = summary_prompt | llm | StrOutputParser()

    summaries = []
    for chunk in document_chunks:
        summary = process_with_retry(summary_chain.invoke, {"doc_chunk": chunk})
        summaries.append(summary)

    combined_summary = " ".join(summaries)

    final_summary_prompt = ChatPromptTemplate.from_template("""
    Please provide a concise overall summary of the following text:

    {combined_summary}

    Overall Summary:
    """)

    final_summary_chain = final_summary_prompt | llm | StrOutputParser()
    return process_with_retry(final_summary_chain.invoke, {"combined_summary": combined_summary})

def query_documents(question, doc_chunks, llm):
    """Query documents with user questions"""
    question_template = ChatPromptTemplate.from_template("""
    Based on the following document chunk, provide a detailed answer to the question. 
    If the information is not available in the chunk, say so.

    Document chunk:
    {chunk}

    Question: {question}

    Answer:
    """)

    question_chain = question_template | llm | StrOutputParser()

    answers = []
    for chunk in doc_chunks:
        answer = process_with_retry(question_chain.invoke, {"chunk": chunk, "question": question})
        answers.append(answer)

    combined_answer_prompt = ChatPromptTemplate.from_template("""
    Combine the following answers into a coherent response:

    {answers}

    Combined answer:
    """)

    combined_answer_chain = combined_answer_prompt | llm | StrOutputParser()
    return process_with_retry(combined_answer_chain.invoke, {"answers": " ".join(answers)})

def main():
    st.set_page_config(page_title='Document Query System', page_icon='📄', layout="wide")
    
    # Initialize Firebase
    db = init_firebase_admin()
    
    # Authentication
    if 'user' not in st.session_state:
        st.title("Document Query System - Login")
        
        tab1, tab2 = st.tabs(["Login", "Sign Up"])
        
        with tab1:
            email = st.text_input("Email", key="login_email")
            password = st.text_input("Password", type="password", key="login_password")
            
            if st.button("Login"):
                try:
                    user = auth.get_user_by_email(email)
                    st.session_state['user'] = {'localId': user.uid, 'email': email}
                    st.success("Logged in successfully!")
                    st.rerun()
                except Exception as e:
                    st.error("Login failed")
        
        with tab2:
            email = st.text_input("Email", key="signup_email")
            password = st.text_input("Password", type="password", key="signup_password")
            organization = st.text_input("Organization Name")
            
            if st.button("Sign Up"):
                try:
                    user = auth.create_user(email=email, password=password)
                    # Create user profile
                    db.collection('users').document(user.uid).set({
                        'email': email,
                        'organization': organization,
                        'created_at': firestore.SERVER_TIMESTAMP
                    })
                    st.success("Account created! Please login.")
                except Exception as e:
                    st.error(f"Sign up failed: {str(e)}")
        return

    # Initialize session state
    for key in ['memory', 'history', 'uploaded_documents', 'combined_document_text']:
        if key not in st.session_state:
            st.session_state[key] = [] if key != 'memory' else ConversationBufferMemory(return_messages=True)

    # User info
    user_id = st.session_state.user['localId']
    user_ref = db.collection('users').document(user_id)
    user_data = user_ref.get().to_dict()

    # Sidebar
    with st.sidebar:
        st.title("User Dashboard")
        st.write(f"Organization: {user_data.get('organization', 'Not Set')}")
        st.write(f"Email: {user_data.get('email')}")
        
        if st.button("Logout"):
            del st.session_state['user']
            st.rerun()
        
        # History
        st.title("History")
        for i, message in enumerate(st.session_state.history):
            st.text_area(f"Entry {i+1}", value=message, height=100, key=f"history_{i}")

    # Main content
    st.title("Document Query System")

    # Document upload section
    st.header("Upload Documents")
    input_method = st.radio("Choose input method:", ("Upload Files", "Paste Text"))

    if input_method == "Upload Files":
        uploaded_files = st.file_uploader(
            "Choose files",
            accept_multiple_files=True,
            type=['pdf', 'docx', 'txt', 'png', 'jpg', 'jpeg']
        )
        if uploaded_files:
            st.session_state.uploaded_documents = []
            st.session_state.combined_document_text = []
            
            for file in uploaded_files:
                with st.spinner(f'Processing {file.name}...'):
                    # Upload to Firebase Storage
                    file_url = upload_file_to_firebase(user_id, file)
                    
                    # Process document
                    text = extract_text_from_file(file)
                    text_chunks = chunk_text(text)
                    thumbnail = generate_thumbnail(file)
                    
                    # Save document info to Firestore
                    doc_ref = user_ref.collection('documents').add({
                        'filename': file.name,
                        'file_url': file_url,
                        'uploaded_at': firestore.SERVER_TIMESTAMP,
                        'type': file.type,
                        'text_content': text,
                        'chunk_count': len(text_chunks)
                    })
                    
                    # Update session state
                    st.session_state.uploaded_documents.append({
                        'file': file,
                        'text_chunks': text_chunks,
                        'thumbnail': thumbnail
                    })
                    st.session_state.combined_document_text.extend(text_chunks)
            
            st.success(f"{len(uploaded_files)} documents processed!")
    else:
        pasted_text = st.text_area("Paste your text here:", height=200)
        if pasted_text:
            text_chunks = chunk_text(pasted_text)
            st.session_state.uploaded_documents = [{
                'file': {'name': 'Pasted Text'},
                'text_chunks': text_chunks,
                'thumbnail': None
            }]
            st.session_state.combined_document_text = text_chunks
            
            # Save to Firestore
            user_ref.collection('documents').add({
                'type': 'pasted_text',
                'content': pasted_text,
                'created_at': firestore.SERVER_TIMESTAMP
            })
            
            st.success("Text processed!")

    # Display documents
    if st.session_state.uploaded_documents:
        st.header("Processed Documents")
        for doc in st.session_state.uploaded_documents:
            col1, col2 = st.columns([1, 3])
            with col1:
                if doc['thumbnail']:
                    st.image(doc['thumbnail'], caption=doc['file'].name)
            with col2:
                st.write(f"Document: {doc['file'].name}")
                st.write(f"Chunks: {len(doc['text_chunks'])}")

        # Analysis section
        st.header("Document Analysis")
        llm = get_llm()

        if st.button("Generate Summary"):
            with st.spinner('Generating summary...'):
                try:
                    summary = generate_document_summary(st.session_state.combined_document_text, llm)
                    st.session_state.history.append(f"Summary: {summary}")
                    
                    # Save to Firestore
                    user_ref.collection('summaries').add({
                        'summary': summary,
                        'created_at': firestore.SERVER_TIMESTAMP
                    })
                    
                    st.write(summary)
                except Exception as e:
                    st.error(f"Error: {str(e)}")

        # Q&A section
        st.header("Ask Questions")
        with st.form(key='question_form'):
            question = st.text_input("Ask about the documents:")
            submit_button = st.form_submit_button(label='Search')

        if submit_button and question:
            with st.spinner('Searching...'):
                try:
                    answer = query_documents(question, st.session_state.combined_document_text, llm)
                    st.session_state.history.append(f"Q: {question}\nA: {answer}")
                    
                    # Save to Firestore
                    user_ref.collection('queries').add({
                        'question': question,
                        'answer': answer,
                        'timestamp': firestore.SERVER_TIMESTAMP
                    })
                    
                    st.write(answer)
                except Exception as e:
                    st.error(f"Error: {str(e)}")

        # Download options
        if st.session_state.history:
            col1, col2 = st.columns(2)
            with col1:
                history_text = "\n\n".join(st.session_state.history)
                bytes_io = io.BytesIO(history_text.encode())
                st.download_button(
                    label="Download History",
                    data=bytes_io,
                    file_name="conversation_history.txt",
                    mime="text/plain"
                )
            
            with col2:
                # Export as PDF option
                if st.button("Export as PDF"):
                    try:
                        pdf = FPDF()
                        pdf.add_page()
                        pdf.set_font("Arial", size=12)
                        
                        # Add organization info
                        pdf.set_font("Arial", 'B', 16)
                        pdf.cell(200, 10, txt=user_data.get('organization', 'Document Analysis Report'), ln=True, align='C')
                        pdf.line(10, 30, 200, 30)
                        
                        # Add date
                        pdf.set_font("Arial", size=10)
                        pdf.cell(200, 10, txt=f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ln=True, align='R')
                        
                        # Add content
                        pdf.set_font("Arial", size=12)
                        for entry in st.session_state.history:
                            pdf.multi_cell(0, 10, txt=entry)
                            pdf.ln()
                        
                        # Save PDF
                        pdf_output = BytesIO()
                        pdf.output(pdf_output)
                        st.download_button(
                            label="Download PDF Report",
                            data=pdf_output.getvalue(),
                            file_name=f"document_analysis_{datetime.now().strftime('%Y%m%d')}.pdf",
                            mime="application/pdf"
                        )
                    except Exception as e:
                        st.error(f"Error generating PDF: {str(e)}")

if __name__ == "__main__":
    main()